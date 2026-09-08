from pathlib import Path
from typing import Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def add_paragraph(
    document,
    text="",
    bold=False,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
):
    paragraph = document.add_paragraph()
    paragraph.alignment = alignment

    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(12)

    return paragraph


def add_numbered_paragraph(document, number, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    number_run = paragraph.add_run(f"{number}. ")
    number_run.bold = True
    number_run.font.size = Pt(12)

    text_run = paragraph.add_run(text)
    text_run.font.size = Pt(12)

    return paragraph


def format_communication_date(date_text):
    """
    Convert:
        15 July 2026
    into:
        15th July 2026
    """

    parts = date_text.strip().split()

    if len(parts) != 3:
        return date_text

    day = int(parts[0])
    month = parts[1]
    year = parts[2]

    if 10 <= day % 100 <= 20:
        suffix = "th"
    else:
        suffix = {
            1: "st",
            2: "nd",
            3: "rd",
        }.get(day % 10, "th")

    return f"{day}{suffix} {month} {year}"


def format_attestation_date(date_text):
    """
    Convert:
        5 September 2026
    into:
        5th day of September 2026
    """

    parts = date_text.strip().split()

    if len(parts) != 3:
        return date_text

    day = int(parts[0])
    month = parts[1]
    year = parts[2]

    if 10 <= day % 100 <= 20:
        suffix = "th"
    else:
        suffix = {
            1: "st",
            2: "nd",
            3: "rd",
        }.get(day % 10, "th")

    return f"{day}{suffix} day of {month} {year}"


def validate_before_generation(mapped_data: Dict):
    """
    Validate the structured case data before generating
    the affidavit.
    """

    errors = []

    document = mapped_data.get("document", {})
    parties = mapped_data.get("parties", {})
    deponent = mapped_data.get("deponent", {})
    reply = mapped_data.get("reply", {})
    attestation = mapped_data.get("attestation", {})
    advocate = mapped_data.get("advocate", {})

    required_fields = [
        ("Court", document.get("court")),
        ("Jurisdiction", document.get("jurisdiction")),
        ("Proceeding Type", document.get("proceeding_type")),
        ("Case Number", document.get("case_number")),
        ("Year", document.get("year")),
        ("Petitioner", parties.get("petitioner")),
        ("Respondent No.1", parties.get("respondent_1")),
        ("Respondent No.2", parties.get("respondent_2")),
        ("Deponent Name", deponent.get("name")),
        ("Deponent Designation", deponent.get("designation")),
        ("Organisation", deponent.get("organisation")),
        ("Address", deponent.get("address")),
        ("Attestation Place", attestation.get("place")),
        ("Attestation Date", attestation.get("date")),
        ("Verification Verb", attestation.get("verification_verb")),
        ("Advocate Firm", advocate.get("firm")),
    ]

    for field_name, value in required_fields:
        if not value:
            errors.append(f"Missing required field: {field_name}")

    points = reply.get("points", {})

    for number in range(1, 7):
        if not points.get(number):
            errors.append(f"Missing Reply Point {number}")

    if not reply.get("communication_date"):
        errors.append("Communication date is missing.")

    if not reply.get("exhibit"):
        errors.append("Exhibit reference is missing.")

    if errors:
        raise ValueError(
            "Pre-generation validation failed:\n"
            + "\n".join(f"- {error}" for error in errors)
        )

    return True


def extract_docx_text(docx_path: str) -> str:
    """
    Extract plain text from the generated DOCX file.
    This text is used by the evaluation module.
    """

    document = Document(docx_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    return "\n".join(paragraphs)


def save_generated_text(docx_path: str, text_path: str):
    """
    Save the generated affidavit as plain text so that
    the evaluator can inspect it.
    """

    text = extract_docx_text(docx_path)

    output = Path(text_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    with open(output, "w", encoding="utf-8") as file:
        file.write(text)

    return output


def generate_affidavit(mapped_data: Dict, output_path: str):
    """
    Generate the Affidavit in Reply using the supplied
    reference structure and case-specific information.

    The function creates:
    1. DOCX affidavit
    2. TXT version for deterministic evaluation
    """

    validate_before_generation(mapped_data)

    document = Document()

    document_data = mapped_data["document"]
    parties = mapped_data["parties"]
    deponent = mapped_data["deponent"]
    reply = mapped_data["reply"]
    attestation = mapped_data["attestation"]
    advocate = mapped_data["advocate"]

    proceeding_type = document_data["proceeding_type"].title()

    communication_date = format_communication_date(
        reply["communication_date"]
    )

    attestation_date = format_attestation_date(
        attestation["date"]
    )

    # ---------------------------------------------------------
    # 1. FORUM / JURISDICTION / CASE NUMBER
    # ---------------------------------------------------------

    add_paragraph(
        document,
        document_data["court"],
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_paragraph(
        document,
        document_data["jurisdiction"],
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    case_number_text = (
        f'{proceeding_type.upper()} NO. '
        f'{document_data["case_number"]} OF '
        f'{document_data["year"]}'
    )

    add_paragraph(
        document,
        case_number_text,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    document.add_paragraph()

    # ---------------------------------------------------------
    # 2. CAUSE TITLE
    # ---------------------------------------------------------

    add_paragraph(
        document,
        parties["petitioner"],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_paragraph(
        document,
        "...Petitioner",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    add_paragraph(
        document,
        "VERSUS",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_paragraph(
        document,
        f"1. {parties['respondent_1']}",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_paragraph(
        document,
        "...Respondent No.1",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    add_paragraph(
        document,
        f"2. {parties['respondent_2']}",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_paragraph(
        document,
        "...Respondent No.2",
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    document.add_paragraph()

    # ---------------------------------------------------------
    # 3. AFFIDAVIT TITLE
    # ---------------------------------------------------------

    add_paragraph(
        document,
        "AFFIDAVIT IN REPLY ON BEHALF OF RESPONDENT NO. 2",
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    document.add_paragraph()

    # ---------------------------------------------------------
    # 4. DEPONENT CLAUSE
    # ---------------------------------------------------------

    deponent_clause = (
        f'I, {deponent["name"]}, '
        f'{deponent["designation"]}, '
        f'having office at {deponent["address"]}, '
        f'the Respondent No.2 above named, '
        f'do hereby solemnly affirm and state as under:'
    )

    add_paragraph(
        document,
        deponent_clause,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    document.add_paragraph()

    # ---------------------------------------------------------
    # 5. REPLY PARAGRAPHS
    # ---------------------------------------------------------

    body_paragraphs = []

    # Point 1
    paragraph_1 = (
        f'I say that I am the Respondent No.2 in the above '
        f'{proceeding_type} and am well acquainted with the facts '
        f'and circumstances of the case. I have perused the '
        f'Writ Petition filed by {parties["petitioner"]} and the '
        f'documents annexed thereto and am competent to affirm '
        f'this Affidavit in Reply. I am filing this Affidavit in '
        f'Reply on behalf of Respondent No.2, '
        f'{deponent["organisation"]}, to oppose the contentions '
        f'raised in the Writ Petition and the reliefs sought by '
        f'the Petitioner.'
    )

    body_paragraphs.append(paragraph_1)

    # Point 2
    paragraph_2 = (
        f'At the outset, I deny each and every allegation, '
        f'contention and submission made in the {proceeding_type}, '
        f'save and except those specifically admitted herein. '
        f'Nothing contained in the Writ Petition that has not '
        f'been specifically dealt with or admitted is to be '
        f'treated as an admission by Respondent No.2.'
    )

    body_paragraphs.append(paragraph_2)

    # Point 3
    paragraph_3 = (
        f'I say that the Writ Petition is misconceived and '
        f'devoid of merits. The actions challenged by the '
        f'Petitioner were taken in accordance with the applicable '
        f'redevelopment procedure and within the authority '
        f'available to Respondent No.2. No legal, constitutional '
        f'or fundamental right of the Petitioner has been infringed.'
    )

    body_paragraphs.append(paragraph_3)

    # Point 4
    paragraph_4 = (
        f'With reference to the specific allegation that the '
        f'impugned communication dated {communication_date} was '
        f'issued without authority, I say that the same is false, '
        f'incorrect and denied.'
    )

    body_paragraphs.append(paragraph_4)

    # Point 5
    paragraph_5 = (
        f'With reference to the impugned communication dated '
        f'{communication_date}, I say that the same was issued '
        f'pursuant to the applicable redevelopment procedure and '
        f'after consideration of the relevant records.'
    )

    body_paragraphs.append(paragraph_5)

    # Point 6
    paragraph_6 = (
        f'I say that Respondent No.2 relies upon the communication '
        f'dated {communication_date}. Hereto annexed and marked '
        f'as EXHIBIT-‘A’ is a copy of the communication dated '
        f'{communication_date}.'
    )

    body_paragraphs.append(paragraph_6)

    # Closing paragraph
    paragraph_7 = (
        f'In the premises aforesaid, I say that the '
        f'{proceeding_type} deserves to be dismissed with costs.'
    )

    body_paragraphs.append(paragraph_7)

    # Add numbered paragraphs
    for number, paragraph_text in enumerate(
        body_paragraphs,
        start=1,
    ):
        add_numbered_paragraph(
            document,
            number,
            paragraph_text,
        )

    body_paragraph_count = len(body_paragraphs)

    document.add_paragraph()

    # ---------------------------------------------------------
    # 6. PRAYER
    # ---------------------------------------------------------

    add_paragraph(
        document,
        "PRAYER",
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    add_paragraph(
        document,
        "I therefore respectfully pray that this Hon’ble Court "
        "may be pleased to:",
    )

    prayer_items = [
        f'dismiss the present {proceeding_type} with costs;',
        "refuse any interim or ad-interim relief sought by the Petitioner; and",
        "grant such other and further reliefs as this Hon’ble Court may deem fit and proper in the facts and circumstances of the case.",
    ]

    for letter, item in zip(
        ["(a)", "(b)", "(c)"],
        prayer_items,
    ):
        paragraph = document.add_paragraph()

        letter_run = paragraph.add_run(f"{letter} ")
        letter_run.bold = True
        letter_run.font.size = Pt(12)

        text_run = paragraph.add_run(item)
        text_run.font.size = Pt(12)

    document.add_paragraph()

    # ---------------------------------------------------------
    # 7. JURAT
    # ---------------------------------------------------------

    verification_verb = attestation.get(
        "verification_verb",
        "",
    ).lower()

    if "swear" in verification_verb:
        jurat_verb = "Sworn"
    else:
        jurat_verb = "Solemnly affirmed"

    add_paragraph(
        document,
        f'{jurat_verb} at {attestation["place"]}',
    )

    add_paragraph(
        document,
        f'On this {attestation_date}',
    )

    document.add_paragraph()

    add_paragraph(
        document,
        "Before Me",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    add_paragraph(
        document,
        "DEPONENT",
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    document.add_paragraph()

    # ---------------------------------------------------------
    # 8. VERIFICATION
    # ---------------------------------------------------------

    add_paragraph(
        document,
        "VERIFICATION",
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )

    verification_text = (
        f'I, {deponent["name"]}, the Deponent above named, '
        f'do hereby verify that the contents of paragraphs 1 to '
        f'{body_paragraph_count} and the Prayer above are true '
        f'and correct to my knowledge and belief and that nothing '
        f'material has been concealed therefrom.'
    )

    add_paragraph(
        document,
        verification_text,
    )

    add_paragraph(
        document,
        f'Verified at {attestation["place"]} on this '
        f'{attestation_date}.',
    )

    add_paragraph(
        document,
        "DEPONENT",
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    document.add_paragraph()

    # ---------------------------------------------------------
    # 9. ADVOCATE BLOCK
    # ---------------------------------------------------------

    add_paragraph(
        document,
        advocate["firm"],
        bold=True,
    )

    add_paragraph(
        document,
        f'Advocates for the {advocate["acting_for"]}',
    )

    # ---------------------------------------------------------
    # SAVE DOCX
    # ---------------------------------------------------------

    output = Path(output_path)

    output.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    document.save(output)

    # ---------------------------------------------------------
    # SAVE TEXT VERSION FOR EVALUATION
    # ---------------------------------------------------------

    text_output = output.with_suffix(".txt")

    save_generated_text(
        str(output),
        str(text_output),
    )

    return output


if __name__ == "__main__":

    from document_parser import extract_text_from_pdf
    from entity_extractor import extract_case_entities
    from mapper import map_case_to_affidavit

    case_path = "data/03_Case_Information.pdf"

    output_path = (
        "outputs/generated_affidavit.docx"
    )

    text = extract_text_from_pdf(
        case_path
    )

    entities = extract_case_entities(
        text
    )

    mapped_data = map_case_to_affidavit(
        entities
    )

    validate_before_generation(
        mapped_data
    )

    output = generate_affidavit(
        mapped_data,
        output_path,
    )

    print(
        "Affidavit generated successfully!"
    )

    print(
        f"DOCX Output: {output}"
    )

    print(
        f"TXT Output: {output.with_suffix('.txt')}"
    )